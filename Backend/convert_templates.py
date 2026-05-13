import os
import re
from docx import Document

# Este script convierte [VARIABLE] a {{ VARIABLE }} en los documentos de Word

def replace_text_in_runs(paragraph):
    # A veces el texto [VAR] está en un solo run, a veces está dividido.
    # Es más fácil juntar todo el texto, buscar y luego reescribir los runs.
    # Pero reconstruir runs pierde formato.
    # Como es una tarea común, una solución sencilla es simplemente hacer
    # un find and replace sobre el texto completo del párrafo. Si no hay formatos 
    # parciales dentro de la variable, podemos limpiar los runs y poner todo en el primero.
    full_text = paragraph.text
    if '[' in full_text and ']' in full_text:
        # Encontramos variables. Reemplazamos [VAR] por {{ VAR }}
        new_text = re.sub(r'\[([^\]]+)\]', r'{{ \1 }}', full_text)
        # Para evitar perder el formato de todo el párrafo, si tiene varios runs, 
        # mantenemos el formato del primer run y borramos los demás (esto asume que 
        # el párrafo tiene formato uniforme. Si no, docxtpl aún funciona con runs unificados)
        if paragraph.runs:
            first_run = paragraph.runs[0]
            first_run.text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""

def process_doc(path, out_path):
    doc = Document(path)
    for p in doc.paragraphs:
        replace_text_in_runs(p)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_text_in_runs(p)
                    
    doc.save(out_path)

if __name__ == "__main__":
    base_dir = r"c:\Users\USUARIO\Documents\Programas Python\Proyecto\Compresores\MinuDoc-Word-Minutes-Generator\EmpresaDeCobros"
    in_pagare = os.path.join(base_dir, "Documentos de apoyo", "PAGARÉ.docx")
    out_pagare = os.path.join(base_dir, "Backend", "apps", "prestamos", "templates", "pagare_template.docx")
    
    in_mutuo = os.path.join(base_dir, "Documentos de apoyo", "CONTRATO DE MUTUO CON GARANTÍA PRENDARIA E HIPOTECARIA.docx")
    out_mutuo = os.path.join(base_dir, "Backend", "apps", "prestamos", "templates", "contrato_template.docx")
    
    os.makedirs(os.path.dirname(out_pagare), exist_ok=True)
    
    process_doc(in_pagare, out_pagare)
    process_doc(in_mutuo, out_mutuo)
    print("Plantillas convertidas exitosamente.")
